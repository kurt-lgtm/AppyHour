"""ONE carrier-invoice reader per carrier — `appyhour_lib.invoice_parsers` (INVOICE_INGEST_RULES
§5.2/§5.3, BUG_LOG 2026-09-07 TWO-READERS).

Fixtures under `tests/fixtures/invoices/` are the REAL rows of the six trackings on which the
local and cloud readers disagreed (names/streets redacted, every parsed field verbatim). Each
test pins the CANONICAL value — the invoice's own arithmetic — not either reader's old answer.
"""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from appyhour_lib import invoice_parsers as ip  # noqa: E402

FIX = Path(__file__).parent / "fixtures" / "invoices"


def _b(name: str) -> bytes:
    return (FIX / name).read_bytes()


# ───────────────────────────────────────────────── the six disagreeing rows, canonical values

def test_ups_multiline_tracking_is_ONE_row_with_its_charge_lines_summed():
    """`1Z2H94940334864194`: freight $18.08 (`Ground Residential`) + a $1.40 `Shipping Charge
    Corrections` line, same invoice. Local kept the LAST line ($1.40, `Ground`) — a 93%
    understatement; cloud wrote two shipments. The invoice bills $19.48 for that package."""
    rows = ip.parse_ups_csv_bytes(_b("AHB_00356_UPS Shipping Breakdown_AHB_6-1-26 (multiline).csv"),
                                  "AHB_00356")
    assert len(rows) == 1
    r = rows[0]
    assert r["tracking"] == "1Z2H94940334864194"
    assert r["cost"] == pytest.approx(19.48)
    assert r["charge_lines"] == 2
    assert r["service"] == "Ground Residential", "the FREIGHT line seeds the descriptive fields"
    assert r["invoice_id"] == "AHB_00356"
    assert r["billed_weight"] == 13.0 and r["weight"] is None, "§5.1: chargeable -> billed_weight"
    assert r["zip_code"] == "87535" and r["state"] == "NM"


@pytest.mark.parametrize("fixture,tracking,outbound,rts,total", [
    ("AHB_00215_OnTrac Shipping Breakdown_AHB_1-5-26 (rts).csv", "1LSD8S9000MXDVF", 8.80, 7.32, 16.12),
    ("AHB_00227_OnTrac Shipping Breakdown_AHB_1-19-26 (rts).csv", "1LSD8S9000NWSU1", 9.53, 6.17, 15.70),
    ("AHB_00194_OnTrac Shipping Breakdown_AHB_12-8-25 (rts).csv", "1LSD8S9000J2Q6S", 6.76, 6.94, 13.70),
    ("AHB_00266_OnTrac Shipping Breakdown_AHB_3-9-26 (rts).csv", "1LSDBVC000R8IRL", 6.30, 5.70, 12.00),
    ("AHB_00266_OnTrac Shipping Breakdown_AHB_3-9-26 (rts).csv", "1LSDBVC000R8G0R", 5.31, 4.72, 10.03),
])
def test_ontrac_return_to_sender_leg_is_a_second_charge_folded_into_the_outbound(
        fixture, tracking, outbound, rts, total):
    """Five `RD` trackings billed TWICE on one invoice: the outbound line and a `Return to
    Sender = Y` line (a different Billing Date, same Invoice Number). Local kept the return leg
    (last-wins); cloud kept both rows. We paid both: cost = outbound + return, flagged RTS."""
    rows = [r for r in ip.parse_ontrac_csv_bytes(_b(fixture), fixture[:9]) if r["tracking"] == tracking]
    assert len(rows) == 1
    r = rows[0]
    assert r["cost"] == pytest.approx(total), f"expected {outbound} + {rts}"
    assert r["charge_lines"] == 2
    assert r["return_to_sender"] is True
    assert r["service"] == "RD"


def test_fedex_rebill_on_a_later_invoice_stays_TWO_rows_one_per_file():
    """`396455860812`: $12.68 on invoice 910790778 (`AHB_00202`), then re-billed $10.14 on
    invoice 911747781 (`AHB_00205`, with POD). These are two invoices, so two rows — the cloud
    had them right. The local sqlite row said `AHB_00202` but carried $10.14: its
    `ON CONFLICT(tracking) DO UPDATE cost=excluded.cost` overwrote the cost and kept the OLD
    invoice_id/source_file — a provenance clobber at the STORE, not a parser difference."""
    a = ip.parse_fedex_xlsx_bytes(_b("AHB_00202_AHB FedEx Breakdown_12-8-25 (sample).XLSX"), "",
                                  source_name="AHB_00202_AHB FedEx Breakdown_12-8-25 (sample).XLSX")
    b = ip.parse_fedex_xlsx_bytes(_b("AHB_00205_AHB FedEx Breakdown_12-15-25 (sample).XLSX"), "",
                                  source_name="AHB_00205_AHB FedEx Breakdown_12-15-25 (sample).XLSX")
    assert [r["cost"] for r in a] == [pytest.approx(12.68)]
    assert [r["cost"] for r in b] == [pytest.approx(10.14)]
    assert a[0]["invoice_id"] == "AHB_00202" and b[0]["invoice_id"] == "AHB_00205", \
        "identity = the AHB_<num> filename token (the same derivation the local scanner uses)"
    assert a[0]["carrier_invoice_no"] == "910790778" and b[0]["carrier_invoice_no"] == "911747781"
    assert a[0]["delivery_date"] is None and b[0]["delivery_date"] == "2025-12-15"
    assert a[0]["ship_date"] == "2025-12-09", "8-digit `20251209` normalized to ISO at the boundary"
    assert a[0]["acct"] == "206137911" and a[0]["hub"] == "Dallas"
    assert a[0]["billed_weight"] == 6.0 and b[0]["billed_weight"] == 7.0


# ───────────────────────────────────────────────── both UPS layouts, by header NAME

def test_both_ups_headed_layouts_parse_and_differ_only_by_the_amount_due_column():
    """Layout A = our account export (`Invoice_000000C411H4###`, carries `Amount Due` on an
    invoice-header row with no tracking); layout B = RMFG's `UPS Shipping Breakdown` (no
    `Amount Due`, no header row). Same reader, resolved by NAME — never by position."""
    a = ip.parse_ups_csv_bytes(_b("Invoice_000000C411H4336_081526.csv"), "000000C411H4336")
    b = ip.parse_ups_csv_bytes(_b("AHB_00356_UPS Shipping Breakdown_AHB_6-1-26.csv"), "AHB_00356")
    assert len(a) == 6 and len(b) == 4
    assert sum(r["cost"] for r in a) == pytest.approx(92.39), "== the file's own Amount Due"
    assert all(r["billed_weight"] and r["weight"] is None for r in a + b)
    assert all(len(r["state"]) == 2 and len(r["zip_code"]) == 5 for r in a + b)
    assert {r["invoice_id"] for r in a} == {"000000C411H4336"}
    assert {r["invoice_id"] for r in b} == {"AHB_00356"}


def test_a_ups_file_missing_a_required_column_RAISES_and_names_it():
    """§4/§2: a header without `Tracking Number` is a dialect we do not have — RAISE, naming the
    missing columns; never `continue` into an empty load (the `if len(r) < 53` no-op)."""
    text = _b("Invoice_000000C411H4336_081526.csv").decode("latin-1")
    broken = text.replace("Tracking Number", "Tracking Nbr", 1)
    with pytest.raises(ip.InvoiceDialectError) as ei:
        ip.parse_ups_csv_bytes(broken.encode("latin-1"), "x")
    assert "Tracking Number" in str(ei.value)
    with pytest.raises(ip.InvoiceDialectError):
        ip.parse_ups_csv_bytes(b"some,unrelated,header\n1,2,3\n", "x")


def test_carrier_is_the_filename_TOKEN_a_reader_fed_the_wrong_carrier_REFUSES():
    """Tracking Coordinator (4): carrier selection is by filename token ONLY; header sniffing is
    for the dialect WITHIN a carrier. OnTrac CSVs also carry `Tracking Number`, so a header-sniff
    carrier guess sends them to the UPS reader — 14 rows came back at $0.00 that way, and the
    reverse (AHB_00350 UPS file → OnTrac reader) wrote 61 `1Z` rows as OnTrac at $0. Neither
    reader may quietly produce rows for the other carrier's file: it must RAISE, naming columns."""
    ontrac = _b("AHB_00215_OnTrac Shipping Breakdown_AHB_1-5-26 (rts).csv")
    ups = _b("AHB_00356_UPS Shipping Breakdown_AHB_6-1-26.csv")
    with pytest.raises(ip.InvoiceDialectError) as ei:
        ip.parse_ups_csv_bytes(ontrac, "AHB_00215")
    assert "Billed Charge" in str(ei.value)
    with pytest.raises(ip.InvoiceDialectError) as ei:
        ip.parse_ontrac_csv_bytes(ups, "AHB_00356")
    assert "Total Charges" in str(ei.value)
    # the right carrier's reader, from the token, still parses both
    assert ip.parse_ontrac_csv_bytes(ontrac, "AHB_00215")[0]["cost"] > 0
    assert ip.parse_ups_csv_bytes(ups, "AHB_00356")[0]["cost"] > 0


def test_ontrac_file_missing_total_charges_RAISES():
    text = _b("AHB_00215_OnTrac Shipping Breakdown_AHB_1-5-26 (rts).csv").decode("latin-1")
    with pytest.raises(ip.InvoiceDialectError) as ei:
        ip.parse_ontrac_csv_bytes(text.replace("Total Charges", "Charges").encode("latin-1"), "x")
    assert "Total Charges" in str(ei.value)


def test_fedex_workbook_skips_sheets_by_NAME_and_raises_when_none_qualifies(capsys):
    """A sheet whose first row lacks `Net Charge Amount` + a tracking column contributes NOTHING
    and is named on stderr; a workbook with no qualifying sheet is an error, never zero rows."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.append(["Some", "Banner", "Row"])
    ws.append([1, 2, 3])
    ws2 = wb.create_sheet("Data")
    ws2.append(["Express or Ground Tracking ID", "Net Charge Amount", "Shipper City",
                "Shipper State", "Recipient State", "Recipient Zip Code", "Bill to Account Number"])
    ws2.append([123456789012, 9.5, "Garland", "TX", "RI", 2879, 206137911.0])
    buf = io.BytesIO()
    wb.save(buf)
    rows = ip.parse_fedex_xlsx_bytes(buf.getvalue(), "AHB_00999")
    assert [r["tracking"] for r in rows] == ["123456789012"]
    assert rows[0]["zip_code"] == "02879", "a leading zero is data"
    assert rows[0]["acct"] == "206137911", "Excel float id collapses to one spelling"
    assert "SKIP sheet 'Dashboard'" in capsys.readouterr().err

    wb2 = openpyxl.Workbook()
    wb2.active.append(["Tracking", "Cost"])
    wb2.active.append(["x", 1])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    with pytest.raises(ip.InvoiceDialectError):
        ip.parse_fedex_xlsx_bytes(buf2.getvalue(), "AHB_00998")


def test_rmfg_fedex_breakdown_without_a_bill_to_column_falls_back_to_the_umbrella_account():
    """Kurt ruling 2026-06-26: an RMFG breakdown IS the RMFG umbrella account — a declared fact.
    Only when the AHB_ token proves it is an RMFG file; a non-RMFG workbook stays `None`."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Express or Ground Tracking ID", "Net Charge Amount", "Shipper City", "Shipper State",
               "Recipient State", "Recipient Zip Code"])
    ws.append(["123456789012", 9.5, "Garland", "TX", "GA", "30301"])
    buf = io.BytesIO()
    wb.save(buf)
    assert ip.parse_fedex_xlsx_bytes(buf.getvalue(), "AHB_00300")[0]["acct"] == "206137911"
    assert ip.parse_fedex_xlsx_bytes(buf.getvalue(), "")[0]["acct"] is None


# ───────────────────────────────────────────────── one module, every consumer

def test_every_local_consumer_reaches_the_SAME_function_objects():
    """The re-exports are the same objects — a second implementation anywhere is the bug."""
    gp = Path(__file__).resolve().parents[1] / "GelPackCalculator"
    if not (gp / "shipping_invoice_db.py").exists():
        pytest.skip("GelPackCalculator (separate repo) not checked out beside appyhour_lib")
    sys.path.insert(0, str(gp))
    import shipping_invoice_db as sidb  # noqa: E402
    from ShippingReports.parsers import ups as sr_ups  # noqa: E402
    for name in ("parse_ups_csv_bytes", "parse_ontrac_csv_bytes", "parse_fedex_xlsx_bytes",
                 "parse_fedex_csv_bytes", "identify_hub", "parse_date_flexible"):
        assert getattr(sidb, name) is getattr(ip, name), name
    assert sr_ups.parse_ups_csv_bytes is ip.parse_ups_csv_bytes


def test_aggregate_keeps_first_appearance_order_and_or_s_the_rts_flag():
    rows = ip.aggregate_charge_lines([
        {"invoice_id": "A", "tracking": "t1", "cost": 1.0, "service": "Ground"},
        {"invoice_id": "A", "tracking": "t2", "cost": 2.0},
        {"invoice_id": "A", "tracking": "t1", "cost": 0.5, "service": "Adj", "return_to_sender": True,
         "delivery_date": "2026-01-02"},
        {"invoice_id": "B", "tracking": "t1", "cost": 9.0},
    ])
    assert [(r["invoice_id"], r["tracking"], r["cost"], r["charge_lines"]) for r in rows] == [
        ("A", "t1", 1.5, 2), ("A", "t2", 2.0, 1), ("B", "t1", 9.0, 1)]
    assert rows[0]["service"] == "Ground" and rows[0]["return_to_sender"] is True
    assert rows[0]["delivery_date"] == "2026-01-02", "a later line may fill what the first lacked"
