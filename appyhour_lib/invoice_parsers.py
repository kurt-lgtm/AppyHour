"""THE carrier-invoice parsers — ONE reader per carrier, shared by every ingest path.

🔴 Constraints SSOT: `ShipRouting/INVOICE_INGEST_RULES.md` §5.2 / §5.3 — read it before changing
anything here. Every consumer imports THIS module:

  - local Kori/sync path   → `GelPackCalculator/shipping_invoice_db.py` (re-exports these names)
  - local ingest scanner   → `GelPackCalculator/auto_import.py` (via shipping_invoice_db)
  - cloud ingest-worker    → `ShipRouting/server/sync_invoices.py` (via the vendored
                             `server/shipping_invoice_db.py`, which also re-exports these names)
  - legacy ShippingReports → `ShippingReports/parsers/ups.py` (thin Shipment adapter)

WHY THIS MODULE EXISTS (the TWO-READERS burn, BUG_LOG 2026-09-07)
The cloud ran a VENDORED copy of the local parser (ShipRouting `4972f7e`) and the two drifted.
Measured on the same files, the two readers disagreed on cost for the same tracking:

  UPS   `1Z2H94940334864194` in `AHB_00356_UPS Shipping Breakdown_AHB_6-1-26.csv` — the file
        carries TWO lines: freight $18.08 (`Ground Residential`, Outbound) + a $1.40 `Shipping
        Charge Corrections` line. Local (sqlite `ON CONFLICT(tracking) DO UPDATE cost=excluded`)
        kept the LAST line: $1.40, service `Ground` — a 93% understatement that looked plausible.
        Cloud inserted BOTH lines as two shipments — a per-tracking consumer sees an arbitrary one.
  OnTrac `1LSD8S9000MXDVF` (`AHB_00215`) — outbound $8.80 + a `Return to Sender = Y` leg $7.32,
        same invoice number. Local kept $7.32 (the return), cloud kept both rows. Four more RD
        trackings (`AHB_00194`, `AHB_00227`, `AHB_00266` ×2) have the same shape.

Neither reader was right. The invoice's own arithmetic is the authority: the control total the
cloud asserts (§3, `Amount Due` == Σ lines) is a SUM over charge lines, so the cost of a tracking
within one invoice file is the SUM of every charge line carrying that tracking. One row per
(invoice file, tracking); the FIRST line seeds the descriptive fields (service, zone, address —
the freight line comes first in every dialect we have seen); `charge_lines` records how many
lines were folded so the fold is never invisible.

RULES THIS MODULE ENFORCES
- §4 header-driven dialects resolve every column BY NAME. A file whose header lacks a required
  column RAISES `InvoiceDialectError` naming the missing columns — never `continue`, never a
  positional guess. An xlsx sheet lacking the required columns is SKIPPED by name (printed), and
  a workbook where NO sheet qualifies raises.
- §5.1 `weight` is the legacy billed-else-actual COALESCE; `billed_weight` is the carrier's
  chargeable weight VERBATIM; a blank weight is `None`, never `0.0`.
- §1 accounts are read VERBATIM from the row (`Bill to Account Number`); canonicalisation is the
  STORE's job (`acct_canon`), never the parser's — except the one Kurt-declared fallback: an RMFG
  FedEx breakdown lacking the column is the RMFG umbrella `206137911` (ruling 2026-06-26).
- Dates are ISO `YYYY-MM-DD` strings at the boundary (the 8,940-row `20260817` incident).
- ZIPs are TEXT, 5 chars, left-padded when Excel handed us a number (02879 → 2879 → `02879`).
- stdlib only at import; `openpyxl` is imported lazily inside the xlsx reader (appyhour_lib rule).
"""
from __future__ import annotations

import csv
import io
import re
import sys
from datetime import date, datetime
from typing import Optional

__all__ = [
    "InvoiceDialectError", "identify_hub", "parse_date_flexible", "iso_date",
    "parse_ups_csv_bytes", "parse_ontrac_csv_bytes",
    "parse_fedex_xlsx_bytes", "parse_fedex_csv_bytes",
    "aggregate_charge_lines", "require_columns",
]

RMFG_FEDEX_UMBRELLA_ACCT = "206137911"     # Kurt ruling 2026-06-26 — a declared fact, not a guess


class InvoiceDialectError(ValueError):
    """The file does not carry the columns this dialect's reader needs. LOUD by construction —
    the message names the file label and every missing column (§4)."""


# ────────────────────────────────────────────────────────────────────── shared helpers

def identify_hub(ref_field: str = "", shipper_city: str = "",
                 shipper_state: str = "") -> str:
    """Hub for a shipment row — real city names ONLY, never a suffix variant (Kurt 2026-08-31:
    "the ahb suffix variant is a dumb naming scheme. just keep it its real names").

    🔴 ONE copy. This used to live in BOTH `GelPackCalculator/shipping_invoice_db.py` and the
    vendored `ShipRouting/server/shipping_invoice_db.py` under a "KEEP IN SYNC" note — the local
    copy sat stale with no Chicago (opened 7/30) and no Swedesboro (opened 8/06), so every invoice
    from either hub landed 'Unknown' and its lanes never earned proof (727 Unknown rows measured
    cloud-side 2026-08-31). A hub added to `ShipRouting/lib/hubs.py` must be added HERE."""
    ref_lower = (ref_field or "").lower()
    for name in ('nashville', 'dallas', 'anaheim', 'chicago', 'swedesboro', 'indianapolis'):
        if name in ref_lower:
            return name.capitalize()

    city_upper = (shipper_city or "").upper().strip()
    state_upper = (shipper_state or "").upper().strip()
    if city_upper == 'GARLAND' or (state_upper == 'TX' and city_upper != 'WOBURN'):
        return 'Dallas'
    if city_upper in ('NASHVILLE', 'ANTIOCH') or state_upper == 'TN':
        return 'Nashville'
    if city_upper == 'ANAHEIM' or state_upper == 'CA':
        return 'Anaheim'
    # Chicago injection sites (opened 2026-07-30): FedEx Bedford Park, OnTrac Romeoville.
    if city_upper in ('CHICAGO', 'BEDFORD PARK', 'ROMEOVILLE') or state_upper == 'IL':
        return 'Chicago'
    # Swedesboro NJ (opened 2026-08-06): OnTrac LTSC 08085, FedEx ex-Barrington 08007.
    if city_upper in ('SWEDESBORO', 'BARRINGTON') or state_upper == 'NJ':
        return 'Swedesboro'
    if city_upper == 'INDIANAPOLIS' or state_upper == 'IN':
        return 'Indianapolis'
    if city_upper == 'WOBURN' or state_upper == 'MA':
        return 'HQ_IGNORE'
    return 'Unknown'


def parse_date_flexible(val) -> Optional[date]:
    """Date from any invoice cell shape; None when unparseable (never a guessed date)."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    if not s:
        return None
    for fmt in ('%Y%m%d', '%m/%d/%Y', '%m/%d/%Y %H:%M', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def iso_date(val) -> Optional[str]:
    d = parse_date_flexible(val)
    return d.isoformat() if d else None


def _txt(v) -> Optional[str]:
    """Excel returns numeric ids as floats: '203738113.0' and '203738113' are the SAME identity
    (the five-spellings-of-two-accounts bug). Strip the trailing .0 at the boundary."""
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        s = s[:-2]
    return s or None


def _zip5(v) -> Optional[str]:
    """ZIP is TEXT. Excel stores `Recipient Zip Code` as a NUMBER, so 02879 arrives as 2879 and
    the validator then REJECTS the row — 70/3,379 rows in one workbook, every one MA/RI/NJ/CT/NH,
    i.e. precisely the Swedesboro lanes. A leading zero is data."""
    s = _txt(v)
    if not s:
        return None
    s = s.split("-")[0].strip()
    if s.isdigit():
        return s.zfill(5)[:5]
    return s[:5] or None


def _num(v) -> Optional[float]:
    """float or None — a blank is ABSENT, never 0.0 (§5.1)."""
    if v in (None, ""):
        return None
    try:
        return float(str(v).strip().replace('"', "").replace(",", ""))
    except (TypeError, ValueError):
        return None


def _pos(v) -> Optional[float]:
    f = _num(v)
    return f if f is not None and f > 0 else None


def _money(v) -> float:
    """A charge cell. Unparseable -> 0.0 rather than a guess (a fabricated charge is
    indistinguishable from a real one downstream)."""
    f = _num(v)
    return f if f is not None else 0.0


def require_columns(names, required, label: str) -> None:
    """§4: the header must carry every required column BY NAME, or the reader RAISES."""
    have = {str(n).strip() for n in names if n is not None}
    missing = [c for c in required if c not in have]
    if missing:
        raise InvoiceDialectError(
            f"{label}: missing expected column(s) {missing} — header has {sorted(have)[:12]}…; "
            f"a new dialect gets its own reader + fixture (INVOICE_INGEST_RULES §5), never a "
            f"positional guess")


_COST_KEYS = ("cost", "charge_lines", "return_to_sender", "is_freight", "invoice_id", "tracking")


def aggregate_charge_lines(rows: list[dict]) -> list[dict]:
    """ONE shipment per (invoice_id, tracking): cost = Σ charge lines, the FREIGHT line seeds every
    other field, `charge_lines` = how many lines were folded. Order of first appearance is kept.

    Why SUM and not first/last: the invoice's own control total (§3) is Σ lines, so any other
    per-tracking rule cannot reconcile to the invoice. Last-wins produced the $1.40 UPS row and the
    $7.32 return-leg OnTrac rows; first-wins would silently drop every correction and surcharge.
    A credit / negative correction SUBTRACTS — never dropped, never abs()'d (real case:
    `1ZC411H40311605473` = 14.92 − 0.61 = $14.31).

    Descriptive fields (service, zone, address, weights, dates) come from the first line whose
    `is_freight` is true (readers set it False for a UPS `Adjustments & Other Charges/...` section
    and for an OnTrac `Return to Sender = Y` leg — those lines carry thin or altered metadata: an
    address-correction line has blank city/state/zip, zone `000`, a pickup date days later). Until
    a freight line is seen the first line stands in and is REPLACED when one arrives; an
    all-adjustment tracking keeps line 1. `return_to_sender` is OR-ed so an RTS leg folded into
    its outbound stays visible. Every fold is written to stderr so the ingest log carries the
    provenance (`shipments` has no `charge_lines` column)."""
    out: dict[tuple, dict] = {}
    for r in rows:
        key = (r.get("invoice_id"), r["tracking"])
        freight = bool(r.get("is_freight", True))
        if key in out:
            agg = out[key]
            agg["cost"] = round(agg["cost"] + (r.get("cost") or 0.0), 2)
            agg["charge_lines"] += 1
            agg["_lines"].append(r.get("cost") or 0.0)
            if r.get("return_to_sender"):
                agg["return_to_sender"] = True
            if freight and not agg["is_freight"]:
                # the freight line arrived after an adjustment/RTS line: its metadata wins
                for k, v in r.items():
                    if k not in _COST_KEYS:
                        agg[k] = v
                agg["is_freight"] = True
            else:
                # a later line may carry what the seed lacked (e.g. POD on a re-bill)
                for k in ("delivery_date", "transit_days"):
                    if agg.get(k) is None and r.get(k) is not None:
                        agg[k] = r[k]
        else:
            r = dict(r)
            r["cost"] = round(r.get("cost") or 0.0, 2)
            r["charge_lines"] = 1
            r["is_freight"] = freight
            r["_lines"] = [r["cost"]]
            r.setdefault("return_to_sender", False)
            out[key] = r
    result = []
    for agg in out.values():
        lines = agg.pop("_lines")
        agg.pop("is_freight", None)
        if agg["charge_lines"] > 1:
            print(f"[invoice_parsers] {agg.get('carrier')} {agg.get('invoice_id')} {agg['tracking']}: "
                  f"cost {agg['cost']:.2f} = {' + '.join(f'{c:.2f}' for c in lines)} "
                  f"({agg['charge_lines']} lines{', RTS' if agg.get('return_to_sender') else ''})",
                  file=sys.stderr, flush=True)
        result.append(agg)
    return result


def _is_adjustment_section(section: str) -> bool:
    """A UPS `Invoice Section` that is an after-the-fact correction, not the freight charge."""
    s = (section or "").lower()
    return "adjustment" in s or "correction" in s


# ────────────────────────────────────────────────────────────────────── UPS

UPS_HEADED_REQUIRED = ("Tracking Number", "Billed Charge", "Service Level", "Pickup Date",
                       "Receiver State", "Receiver Zip Code")


def parse_ups_csv_bytes(csv_bytes: bytes, invoice_id: str = "") -> list[dict]:
    """UPS — BOTH dialects, auto-detected by SHAPE (never by filename):
      1. headed CSV (`Tracking Number` / `Billed Charge` / `Incentive Credit`; both our
         `Invoice_000000C411H4*` account exports and RMFG's `AHB_*_UPS Shipping Breakdown` use it)
      2. Detailed Billing v2.1 — no header, first field `2.1`, 250+ positional fields.
    Returns ONE row per tracking (charge lines summed)."""
    text = csv_bytes.decode('latin-1')
    first_line = text.split('\n', 1)[0]
    if 'Tracking Number' in first_line:
        return _parse_ups_headed(text, invoice_id)
    if first_line.startswith('2.1,') or first_line.count(',') >= 80:
        return _parse_ups_billing_data(text, invoice_id)
    raise InvoiceDialectError(
        f"UPS {invoice_id or '?'}: neither the headed dialect (`Tracking Number` in row 1) nor "
        f"Detailed Billing v2.1 (row 1 starts `2.1,` / 80+ positional fields) — first line: "
        f"{first_line[:80]!r}")


def _parse_ups_headed(text: str, invoice_id: str) -> list[dict]:
    reader = csv.DictReader(io.StringIO(text))
    require_columns(reader.fieldnames or (), UPS_HEADED_REQUIRED, f"UPS headed {invoice_id or '?'}")
    lines = []
    for row in reader:
        tracking = (row.get('Tracking Number') or '').strip()
        if not tracking:
            continue        # the invoice-header row (`Amount Due` lives there) has no tracking
        pickup = parse_date_flexible(row.get('Pickup Date'))
        # §5 dialect 2: cost = Billed Charge + Incentive Credit (the credit is negative)
        cost = _money(row.get('Billed Charge')) + _money(row.get('Incentive Credit'))
        lines.append({
            "tracking": tracking,
            "carrier": "UPS",
            "service": (row.get('Service Level') or '').strip(),
            "hub": identify_hub(ref_field=row.get('Reference No.2') or '',
                                shipper_city=row.get('Sender City') or '',
                                shipper_state=row.get('Sender State') or ''),
            "state": (row.get('Receiver State') or '').strip().upper(),
            "zip_code": _zip5(row.get('Receiver Zip Code')),
            "city": (row.get('Receiver City') or '').strip(),
            "zone": (row.get('Zone') or '').strip(),
            "cost": cost,
            # 🔴 `Weight` here is integer CHARGEABLE weight -> billed_weight, NOT weight (§5.1)
            "weight": None,
            "billed_weight": _pos(row.get('Weight')),
            "ship_date": pickup.isoformat() if pickup else None,
            "delivery_date": None,
            "transit_days": None,
            "invoice_id": invoice_id,
            "acct": (row.get('Account Number') or '').strip() or None,
            "invoice_section": (row.get('Invoice Section') or '').strip() or None,
            "is_freight": not _is_adjustment_section(row.get('Invoice Section')),
        })
    return aggregate_charge_lines(lines)


def _parse_ups_billing_data(text: str, invoice_id: str) -> list[dict]:
    """UPS Detailed Billing v2.1 (headerless, 250-col). Positional addressing is unavoidable here
    (§4), which makes `appyhour_lib.shipment_validate` the ONLY net downstream.

    Column map (0-indexed): 11 pickup date · 13/20 tracking · 16/22 hub reference · 26 weight ·
    33 zone · 43 line type (FRT freight / ACC accessorial / INF info) · 45 service ·
    52 billed charge · 70/71 sender city/state · 78/79/80 receiver city/state/zip.

    🔴 EVERY line's charge is summed — never FRT-only. `ShippingReports/parsers/ups.py` filtered
    `row[43] == 'FRT'` and dropped every ACC accessorial line (fuel, DAS, residential…): the cloud
    understated UPS v2.1 cost by 9.02% against the local reader (handoff
    `cost-parser-disagreement-2026-09-07`). The FRT line seeds the descriptive fields; an
    accessorial line contributes dollars only."""
    lines = []
    for row in csv.reader(io.StringIO(text)):
        if len(row) < 81:
            if any(c.strip() for c in row):
                # a short row is a dialect drift, not a blank — say so (§2: never silent-continue)
                print(f"[invoice_parsers] UPS v2.1 {invoice_id}: short row ({len(row)} cols) "
                      f"skipped: {','.join(row[:6])!r}", file=sys.stderr, flush=True)
            continue
        tracking = row[20].strip() or row[13].strip()
        if not tracking or not tracking.startswith('1Z'):
            continue
        line_type = row[43].strip().upper()
        pickup = parse_date_flexible(row[11].strip())
        lines.append({
            "is_freight": line_type in ('FRT', ''),
            "tracking": tracking,
            "carrier": "UPS",
            "service": row[45].strip(),
            "hub": identify_hub(ref_field=row[22].strip() or row[16].strip(),
                                shipper_city=row[70].strip(), shipper_state=row[71].strip()),
            "state": row[79].strip().upper(),
            "zip_code": _zip5(row[80].strip()),
            "city": row[78].strip(),
            "zone": row[33].strip(),
            "cost": _money(row[52]),
            "weight": _pos(row[26]),        # legacy coalesce column (this dialect's only weight)
            "billed_weight": None,
            "ship_date": pickup.isoformat() if pickup else None,
            "delivery_date": None,
            "transit_days": None,
            "invoice_id": invoice_id,
            "acct": None,
        })
    return aggregate_charge_lines(lines)


# ────────────────────────────────────────────────────────────────────── OnTrac

ONTRAC_REQUIRED = ("Tracking Number", "Total Charges", "Destination Postalcode",
                   "First Scan Date Time", "Proof of Delivery DateTime")


def parse_ontrac_csv_bytes(csv_bytes: bytes, invoice_id: str = "") -> list[dict]:
    """OnTrac `Shipping Breakdown` CSV (headered). ONE row per tracking; a `Return to Sender = Y`
    line is a second CHARGE for the same package and is summed, with `return_to_sender=True`."""
    text = csv_bytes.decode('latin-1')
    reader = csv.DictReader(io.StringIO(text))
    require_columns(reader.fieldnames or (), ONTRAC_REQUIRED, f"OnTrac {invoice_id or '?'}")
    lines = []
    for row in reader:
        tracking = (row.get('Tracking Number') or '').strip()
        if not tracking:
            continue
        scan_date = parse_date_flexible(row.get('First Scan Date Time'))
        pod_date = parse_date_flexible(row.get('Proof of Delivery DateTime'))
        actual_wt = _pos(row.get("Weight(lbs)"))
        billed_wt = _pos(row.get("Billed Weight (lbs)"))
        lines.append({
            "tracking": tracking,
            "carrier": "OnTrac",
            "service": (row.get('Service Code') or 'RD').strip() or 'RD',
            "hub": identify_hub(ref_field=row.get('Reference1') or ''),
            "state": (row.get('Destination State') or '').strip().upper(),
            "zip_code": _zip5(row.get('Destination Postalcode')),
            "city": (row.get('Destination City') or '').strip(),
            "zone": (row.get('Zone') or '').strip(),
            "cost": _money(row.get('Total Charges')),
            "weight": billed_wt or actual_wt,      # legacy coalesce (§5.1)
            "billed_weight": billed_wt,
            "actual_weight": actual_wt,
            "ship_date": scan_date.isoformat() if scan_date else None,
            "delivery_date": pod_date.isoformat() if pod_date else None,
            "transit_days": (pod_date - scan_date).days if scan_date and pod_date else None,
            "invoice_id": invoice_id,
            "acct": None,
            "return_to_sender": (row.get('Return to Sender') or '').strip().upper() == 'Y',
            "is_freight": (row.get('Return to Sender') or '').strip().upper() != 'Y',
            "dim_l": _pos(row.get("Length(in)")),
            "dim_w": _pos(row.get("Width(in)")),
            "dim_h": _pos(row.get("Height(in)")),
            "dim_factor": _pos(row.get("DIM Factor")),
        })
    return aggregate_charge_lines(lines)


# ────────────────────────────────────────────────────────────────────── FedEx

FEDEX_TRACKING_COLS = ("Express or Ground Tracking ID", "Tracking ID", "Tracking Number")
FEDEX_REQUIRED = ("Net Charge Amount",)
_AHB_TOKEN = re.compile(r"(AHB_\d+)")


def _fedex_row(g, tracking: str, invoice_id: str, file_invoice_token: Optional[str]) -> dict:
    hub = identify_hub(shipper_city=str(g('Shipper City') or ''),
                       shipper_state=str(g('Shipper State') or ''))
    ship_date = parse_date_flexible(g('Shipment Date'))
    pod_date = parse_date_flexible(g('POD Delivery Date') if g('POD Delivery Date') is not None
                                   else g('Delivery Date'))
    rated = _pos(g('Rated Weight Amount (lbs)')) or _pos(g('Rated Weight Amount')) or _pos(g('Rated Weight'))
    actual = (_pos(g('Actual Weight Amount (lbs)')) or _pos(g('Actual Weight Amount'))
              or _pos(g('Actual Weight')))
    raw_acct = _txt(g('Bill to Account Number'))
    # invoice identity: the AHB_<num> filename token when the file has one (ONE derivation, the
    # same one the local scanner uses — two derivations double-ingested $71k, 2026-09-06);
    # otherwise the row's own Invoice Number; otherwise whatever the caller passed.
    inv = file_invoice_token or _txt(g('Invoice Number')) or invoice_id
    return {
        "tracking": tracking,
        "carrier": "FedEx",
        "acct": raw_acct,
        "carrier_invoice_no": _txt(g('Invoice Number')),
        "service": str(g('Service Type') or '').strip(),
        "hub": hub,
        "state": str(g('Recipient State') or '').strip().upper(),
        "zip_code": _zip5(g('Recipient Zip Code')),
        "city": str(g('Recipient City') or '').strip(),
        "zone": _txt(g('Zone Code')) or '',
        "cost": _money(g('Net Charge Amount')),
        "weight": rated or actual,             # legacy coalesce (§5.1)
        "billed_weight": rated,                # chargeable, verbatim
        "actual_weight": actual,
        "ship_date": ship_date.isoformat() if ship_date else None,
        "delivery_date": pod_date.isoformat() if pod_date else None,
        "transit_days": (pod_date - ship_date).days if ship_date and pod_date else None,
        "invoice_id": inv,
    }


def parse_fedex_xlsx_bytes(xlsx_bytes: bytes, invoice_id: str = "", *,
                           include_internal: bool = False,
                           source_name: str = "") -> list[dict]:
    """FedEx workbooks — RMFG `Shipping Breakdown` (one `Data` sheet) AND the `Service Cost
    Analysis` workbooks whose layouts DIFFER (`Details` vs `FedEx Data Clean`, banner rows).

    Every sheet is examined; a sheet is used only if its FIRST row carries `Net Charge Amount`
    and one of the tracking columns BY NAME — others are SKIPPED and named on stderr. A workbook
    with NO qualifying sheet RAISES (§2: never an empty load). `include_internal=False` drops
    HQ_IGNORE (Woburn) shipper rows — the local Kori behaviour; the cloud passes True and lets
    `is_internal` classify."""
    import openpyxl   # lazy — appyhour_lib stays stdlib-only at import
    label = source_name or invoice_id or "FedEx xlsx"
    token = _AHB_TOKEN.search(source_name or invoice_id or "")
    file_token = token.group(1) if token else None
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    lines, used = [], 0
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            hdr = next(it, None)
            names = [str(h).strip() if h is not None else "" for h in (hdr or ())]
            trk_name = next((n for n in FEDEX_TRACKING_COLS if n in names), None)
            missing = [c for c in FEDEX_REQUIRED if c not in names]
            if trk_name is None or missing:
                print(f"[invoice_parsers] {label}: SKIP sheet {ws.title!r} — lacks "
                      f"{missing + ([] if trk_name else ['a tracking column'])} (by NAME, §4)",
                      file=sys.stderr, flush=True)
                continue
            used += 1
            idx = {n: i for i, n in enumerate(names) if n}
            trk = idx[trk_name]
            for r in it:
                if not r or trk >= len(r) or r[trk] in (None, ""):
                    continue

                def g(col, _r=r, _idx=idx):
                    i = _idx.get(col)
                    return _r[i] if i is not None and i < len(_r) else None

                row = _fedex_row(g, _txt(r[trk]), invoice_id, file_token)
                if row["hub"] == 'HQ_IGNORE' and not include_internal:
                    continue
                if row["acct"] is None and file_token:
                    row["acct"] = RMFG_FEDEX_UMBRELLA_ACCT
                lines.append(row)
    finally:
        wb.close()
    if used == 0:
        raise InvoiceDialectError(
            f"{label}: no sheet carries `Net Charge Amount` + a tracking column "
            f"({'/'.join(FEDEX_TRACKING_COLS)}) — sheets: {wb.sheetnames}")
    return aggregate_charge_lines(lines)


def parse_fedex_csv_bytes(csv_bytes: bytes, invoice_id: str = "", *,
                          include_internal: bool = False, source_name: str = "") -> list[dict]:
    """FedEx Billing Online CSV export (headed; one file may hold several invoice numbers).
    Same column names as the xlsx dialect. Per-row `Invoice Number` is the identity unless the
    filename carries an `AHB_<num>` token."""
    text = csv_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    names = reader.fieldnames or ()
    label = source_name or invoice_id or "FedEx csv"
    trk_name = next((n for n in FEDEX_TRACKING_COLS if n in names), None)
    if trk_name is None:
        require_columns(names, ("Express or Ground Tracking ID",), label)
    require_columns(names, FEDEX_REQUIRED, label)
    token = _AHB_TOKEN.search(source_name or "")
    file_token = token.group(1) if token else None
    lines = []
    for r in reader:
        tracking = (r.get(trk_name) or "").strip()
        if not tracking:
            continue
        row = _fedex_row(r.get, tracking, invoice_id, file_token)
        if row["hub"] == 'HQ_IGNORE' and not include_internal:
            continue
        lines.append(row)
    return aggregate_charge_lines(lines)
